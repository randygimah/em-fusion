"""The glossary the scorer reads mid-flow. Spec: eval-009.

A marker meets `track_rate_excess`, `smoothing constant` and `nominal` in
the first pack and will meet them nineteen more times. Stopping to ask what
one means costs the session its rhythm; guessing costs the study a score.

**The risk here is larger than in anything else the pack does, and it is
worth stating before the tests.** A definition that explains what a rule
*implies* rather than what it *denotes* hands the marker the causal link
that evidence use exists to test -- and cause correctness, in pass two,
exists to test again. "An aircraft stopped sending position reports" is a
definition. "...which can indicate interference" is the answer. This is
`eval-007`'s summarising-pack failure in its most dangerous form, because
it is reasoning in prose rather than arithmetic, and prose does not look
like a computation.

So the constraint is enforced by a list of causal connectives no definition
may contain, kept as one named constant with its argument attached, on the
pattern `BLINDING_GIVEAWAYS` already uses.
"""

from __future__ import annotations

import io

import openpyxl
import pytest

from app.ai.service import ConfigurationError
from app.evaluation.pack import (
    BLINDING_GIVEAWAYS,
    CAUSAL_CONNECTIVES,
    GLOSSARY,
    build_pack,
)
from app.evaluation.workbook import TERMS_SHEET, score_sheet_workbook
from tests.test_ai_template import fault_bundle, zone_bundle
from tests.test_eval_pack import one_bundle_campaign, packed

# ------------------------------------------------------- what it must contain


def test_every_citation_key_in_the_pack_is_defined() -> None:
    # spec: eval-009 -- the marker cannot check a citation against the pack
    # without knowing what the cited rule is.
    bundles = [(1, fault_bundle()), (2, zone_bundle())]
    pack = packed(bundles=bundles)
    glossary = pack.hypotheses_md.split("## Evidence pack")[0]

    for _, bundle in bundles:
        for item in bundle.evidence:
            assert item.rule_id in glossary, f"{item.rule_id} undefined"
            assert item.finding_type in glossary, f"{item.finding_type} undefined"


def test_an_undefined_rule_is_refused_rather_than_omitted() -> None:
    # spec: eval-009 -- the failure case that matters. Omitted, a marker
    # meets an undefined rule mid-flow and the pack looks complete; refused,
    # whoever adds a rule has to define it.
    bundle = zone_bundle()
    invented = bundle.with_evidence(
        [
            (
                item
                if item.finding_type != "rf_noise_deviation"
                else _retyped(item, "INVENTED-RULE", "invented_finding")
            )
            for item in bundle.evidence
        ]
    )

    with pytest.raises(ConfigurationError, match="glossary"):
        packed(bundles=[(1, invented)])


def test_only_terms_the_scorer_will_meet_are_listed() -> None:
    # spec: eval-009 -- a definition for something absent from the pack is
    # one more line to read past, and twenty of them is a page of them.
    pack = packed(bundles=[(1, fault_bundle())])
    glossary = pack.hypotheses_md.split("## Evidence pack")[0]

    assert "SENSOR-FAULT" in glossary
    assert (
        "FUSION-SPATIAL" not in glossary
    ), "a fault-only pack defined a rule it never shows"


# ------------------------------------------------ what it must never contain


def test_no_definition_supplies_a_cause() -> None:
    # spec: eval-009 -- the governing constraint, asserted over the whole
    # constant rather than over what one fixture happens to render, because
    # a definition added later is exactly where this would go wrong.
    for term, definition in GLOSSARY.items():
        lowered = definition.lower()
        for connective in CAUSAL_CONNECTIVES:
            assert connective not in lowered, (
                f"the definition of {term!r} says {connective!r}, which tells "
                "the marker what the evidence means rather than what the word does"
            )


def test_the_guard_catches_a_definition_that_explains_too_much() -> None:
    # spec: eval-009 -- the guard is fired at a real violation, because a
    # check verified only against the cases it must pass is verified against
    # the author's model of the world.
    tempting = (
        "one aircraft's position reports stopping, which indicates that "
        "something is interfering with reception in that area"
    )
    assert any(word in tempting for word in CAUSAL_CONNECTIVES)


def test_no_definition_names_an_arm() -> None:
    # spec: eval-009, eval-004 -- a third artefact the marker reads. Left
    # unscanned it does not weaken the blinding, it relocates the leak.
    for term, definition in GLOSSARY.items():
        for giveaway in BLINDING_GIVEAWAYS:
            assert giveaway not in f"{term} {definition}".lower(), term


# ------------------------------------------------------- where it has to be


def test_the_glossary_precedes_the_first_pack() -> None:
    # spec: eval-009 -- after them it is an appendix nobody scrolls to.
    md = packed().hypotheses_md

    assert "## Evidence pack" in md
    assert md.index("Terms used in these packs") < md.index("## Evidence pack")


def test_the_glossary_travels_with_the_workbook_too() -> None:
    # spec: eval-009 -- the workbook is what is open while a score is being
    # entered. A glossary only in the other document is one the scorer has
    # to leave the sheet to reach, which is the interruption this removes.
    pack = packed()
    book = openpyxl.load_workbook(io.BytesIO(score_sheet_workbook(pack.items)))

    assert TERMS_SHEET in book.sheetnames
    listed = {
        str(row[0].value)
        for row in book[TERMS_SHEET].iter_rows(min_row=2)
        if row[0].value
    }
    assert "TRACK-GAP" in listed
    assert listed <= set(GLOSSARY)


def _retyped(item, rule_id: str, finding_type: str):
    from dataclasses import replace

    return replace(item, rule_id=rule_id, finding_type=finding_type)


def test_a_term_used_only_in_an_explanation_is_still_defined() -> None:
    # spec: eval-009 -- the glossary is scanned over everything the marker
    # reads, not only over the findings.
    #
    # `corroborate` is the case that proved it: it appears in the
    # recommended action of an explanation and nowhere in any finding, so a
    # scan over the evidence alone left out one of the words most likely to
    # stop a scorer mid-pack. Pooled across twenty packs and both labels,
    # what the glossary contains says nothing about any single pack.
    hypothesis = _with_action(
        "Corroborate the deviation against a source outside this system."
    )
    rows, _ = one_bundle_campaign(1)
    pack = build_pack(
        generations=rows,
        bundles=[(1, fault_bundle())],
        hypotheses={row.hypothesis_id: hypothesis for row in rows},
        seed=11,
    )
    glossary = pack.hypotheses_md.split("## Evidence pack")[0]

    assert "corroborate" in glossary


def _with_action(action: str):
    from app.ai.hypothesis import Hypothesis
    from app.ai.template import generate

    base = generate(fault_bundle()).model_dump()
    base["recommended_human_action"] = action
    return Hypothesis(**base)
